"""Procurement report generation (PDF + DOCX) from a persisted analysis.

Curated, not a dump. The report leads with a verdict and the handful of strongest,
most-relevant standards the officer should pass to the procurement agency — in plain
language — followed by a compact compliance table and an officer sign-off block.
Low-relevance partial matches stay in the app, not the PDF. Deterministic, no LLM.
"""

from __future__ import annotations

import io
from datetime import datetime, timezone

from sqlalchemy import func, select
from sqlalchemy.orm import Session

from app.models import (
    Analysis, Conflict, CoverageResult, Document, Evidence, Gap, QcoRecord,
    Recommendation, RecommendationEvidence, Requirement, Review, ReviewDecision, Standard,
)

# Max strong matches shown in the report body; the rest are summarised as a count.
_MAX_TOP = 8

# Applicability class → plain-language instruction for the procurement agency.
_CLASS_ADVICE = {
    "DIRECTLY_APPLICABLE": "Require full compliance with this standard in the specification.",
    "NORMATIVE_REFERENCE": "Cite this standard as a normative requirement.",
    "TESTING": "Require test certificates demonstrating conformity to this standard.",
    "SAFETY": "Require safety conformance to this standard.",
    "MATERIAL": "Specify materials conforming to this standard.",
    "INSTALLATION": "Require installation and workmanship as per this standard.",
    "CERTIFICATION": "Require valid certification under this standard.",
    "CONDITIONAL": "Apply this standard where the stated condition holds.",
    "RELATED": "Consider this standard as supporting guidance.",
}


def _gather(db: Session, analysis_id: str) -> dict:
    analysis = db.get(Analysis, analysis_id)
    document = db.get(Document, analysis.document_id) if analysis else None
    reqs = db.execute(select(Requirement).where(Requirement.analysis_id == analysis_id)
                      .order_by(Requirement.req_code)).scalars().all()
    recs = db.execute(select(Recommendation).where(
        Recommendation.analysis_id == analysis_id, Recommendation.excluded == False)  # noqa: E712
        .order_by(Recommendation.requirement_id, Recommendation.final_rank)).scalars().all()
    std_by_id = {s.id: s for s in db.execute(select(Standard)).scalars()}
    decisions = {}
    review = db.execute(select(Review).where(Review.analysis_id == analysis_id)).scalar_one_or_none()
    if review:
        for d in db.execute(select(ReviewDecision).where(ReviewDecision.review_id == review.id)).scalars():
            decisions[(d.target_type, d.target_id)] = d
    has_demo = any(std_by_id[r.standard_id].data_origin == "DEMO_SYNTHETIC" for r in recs)
    return {"analysis": analysis, "document": document, "reqs": reqs, "recs": recs,
            "std_by_id": std_by_id, "decisions": decisions, "has_demo": has_demo}


def _curate(db: Session, analysis_id: str, data: dict) -> dict:
    """Reduce the raw analysis to a curated summary: verdict + top matches + actions."""
    recs, std_by_id = data["recs"], data["std_by_id"]

    # One best-ranked recommendation per distinct standard.
    best_by_std: dict[str, Recommendation] = {}
    for r in recs:
        cur = best_by_std.get(r.standard_id)
        if cur is None or r.relevance_score > cur.relevance_score:
            best_by_std[r.standard_id] = r
    uniq = list(best_by_std.values())

    def is_strong(r: Recommendation) -> bool:
        return bool(r.is_primary) or r.relevance in ("HIGH", "MEDIUM") or r.relevance_score >= 0.5

    strong = sorted([r for r in uniq if is_strong(r)], key=lambda r: r.relevance_score, reverse=True)
    top = strong[:_MAX_TOP]
    appendix_count = len(uniq) - len(top)

    # QCO-mandatory flags for the recommended standards.
    ids = list(best_by_std.keys())
    qco_mandatory: dict[str, QcoRecord] = {}
    if ids:
        for q in db.execute(select(QcoRecord).where(
                QcoRecord.standard_id.in_(ids), QcoRecord.qco_status == "MANDATORY")).scalars():
            qco_mandatory[q.standard_id] = q

    # Coverage + conflicts + gaps for the verdict.
    covs = db.execute(select(CoverageResult.coverage).where(
        CoverageResult.analysis_id == analysis_id)).scalars().all()
    total = len(covs)
    full = sum(1 for c in covs if c == "FULL")
    partial = sum(1 for c in covs if c == "PARTIAL")
    missing = sum(1 for c in covs if c == "MISSING")
    conflicts = db.execute(select(Conflict).where(Conflict.analysis_id == analysis_id)).scalars().all()
    gaps = db.execute(select(func.count()).select_from(Gap).where(Gap.analysis_id == analysis_id)).scalar_one()
    compliance_pct = round(100 * full / total) if total else 0

    if missing > 0 or conflicts:
        verdict = "ACTION REQUIRED"
        verdict_detail = ("This specification should not be tendered until the items below are resolved. "
                          f"{missing} requirement(s) have no applicable standard"
                          f"{f' and {len(conflicts)} conflict(s) were detected' if conflicts else ''}.")
        verdict_color = "#B42318"
    elif partial > 0 or gaps > 0:
        verdict = "REVIEW RECOMMENDED"
        verdict_detail = ("The specification is largely covered but the officer should review the "
                          f"{partial} partially-covered requirement(s) before tendering.")
        verdict_color = "#B54708"
    elif total > 0:
        verdict = "READY TO TENDER"
        verdict_detail = "Every requirement maps to an applicable standard with no unresolved conflicts."
        verdict_color = "#067647"
    else:
        verdict = "REVIEW REQUIRED"
        verdict_detail = "No coverage could be computed — verify the requirements were extracted."
        verdict_color = "#475467"

    # "Advise the procurement agency" — the top 5 concrete actions, most critical first.
    actions: list[str] = []
    for sid, q in qco_mandatory.items():
        s = std_by_id.get(sid)
        if s:
            actions.append(f"MANDATORY: {s.is_number} is QCO-notified — bidders must hold BIS certification"
                           f"{f' under {q.order_name}' if q.order_name else ''}. Reject uncertified offers.")
    for c in conflicts[:2]:
        actions.append(f"Resolve conflict on {c.parameter or 'a parameter'}: "
                       f"{c.value_a}{c.unit_a} vs {c.value_b}{c.unit_b} — confirm the correct value.")
    if missing > 0:
        actions.append(f"{missing} requirement(s) have no applicable standard on file — verify with the "
                       "technical department before issuing the tender.")
    for r in top:
        if len(actions) >= 5:
            break
        s = std_by_id.get(r.standard_id)
        if s and r.standard_id not in qco_mandatory:
            actions.append(f"Include {s.is_number} ({s.title}) as a compliance requirement.")

    top_rows = []
    for r in top:
        s = std_by_id[r.standard_id]
        advice = _CLASS_ADVICE.get(r.applicability_class, "Include as a compliance requirement.")
        top_rows.append({
            "is_number": s.is_number, "title": s.title,
            "relevance": r.relevance, "match_pct": round(r.relevance_score * 100),
            "applicability": r.applicability_class.replace("_", " ").title(),
            "mandatory": r.standard_id in qco_mandatory,
            "advice": advice, "is_demo": s.data_origin == "DEMO_SYNTHETIC",
        })

    return {
        "verdict": verdict, "verdict_detail": verdict_detail, "verdict_color": verdict_color,
        "compliance_pct": compliance_pct, "requirements_total": total,
        "covered": full, "partial": partial, "missing": missing,
        "conflicts": len(conflicts), "gaps": gaps,
        "actions": actions[:5], "top_rows": top_rows, "appendix_count": appendix_count,
        "mandatory_count": len(qco_mandatory),
    }


def build_pdf(db: Session, analysis_id: str, officer=None) -> bytes:
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
    from reportlab.lib import colors

    data = _gather(db, analysis_id)
    c = _curate(db, analysis_id, data)

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, title="MORPHEUS Procurement Compliance Summary",
                            topMargin=18 * mm, bottomMargin=16 * mm, leftMargin=18 * mm, rightMargin=18 * mm)
    styles = getSampleStyleSheet()
    green = colors.HexColor("#0B5D3B")
    h1 = ParagraphStyle("h1", parent=styles["Title"], fontSize=18, textColor=green, spaceAfter=2, alignment=0)
    sub = ParagraphStyle("sub", parent=styles["Normal"], fontSize=9, textColor=colors.HexColor("#667085"))
    body = ParagraphStyle("body", parent=styles["Normal"], fontSize=9.5, leading=13)
    h2 = ParagraphStyle("h2", parent=styles["Heading2"], fontSize=11.5, textColor=colors.HexColor("#111827"),
                        spaceBefore=10, spaceAfter=4)
    story = []

    # ── Header ──
    story.append(Paragraph("Procurement Compliance Summary", h1))
    story.append(Paragraph("MORPHEUS — Standards Intelligence for Public Procurement", sub))
    story.append(Spacer(1, 6))
    doc_name = data["document"].filename if data["document"] else "—"
    meta = f"<b>Tender:</b> {data['analysis'].title or doc_name} &nbsp;&nbsp; " \
           f"<b>Sector:</b> {data['analysis'].sector or '—'} &nbsp;&nbsp; " \
           f"<b>Generated:</b> {datetime.now(timezone.utc).strftime('%d %b %Y, %H:%M UTC')}"
    story.append(Paragraph(meta, sub))
    story.append(Spacer(1, 10))

    # ── Verdict banner ──
    vstyle = ParagraphStyle("v", parent=body, textColor=colors.white, fontSize=12, leading=15)
    vdetail = ParagraphStyle("vd", parent=body, textColor=colors.white, fontSize=9)
    vtbl = Table([[Paragraph(f"<b>{c['verdict']}</b>", vstyle)],
                  [Paragraph(c["verdict_detail"], vdetail)]], colWidths=[doc.width])
    vtbl.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor(c["verdict_color"])),
        ("LEFTPADDING", (0, 0), (-1, -1), 12), ("RIGHTPADDING", (0, 0), (-1, -1), 12),
        ("TOPPADDING", (0, 0), (0, 0), 10), ("BOTTOMPADDING", (0, 1), (0, 1), 10),
    ]))
    story.append(vtbl)
    story.append(Spacer(1, 4))

    # ── Compliance snapshot ──
    snap = [[f"{c['compliance_pct']}%", str(c["requirements_total"]), str(c["covered"]),
             str(c["partial"]), str(c["missing"]), str(c["mandatory_count"])],
            ["Compliance", "Requirements", "Covered", "Partial", "Not covered", "QCO mandatory"]]
    snaptbl = Table(snap, colWidths=[doc.width / 6] * 6)
    snaptbl.setStyle(TableStyle([
        ("FONTSIZE", (0, 0), (-1, 0), 15), ("FONTSIZE", (0, 1), (-1, 1), 7.5),
        ("TEXTCOLOR", (0, 0), (-1, 0), green), ("TEXTCOLOR", (0, 1), (-1, 1), colors.HexColor("#667085")),
        ("ALIGN", (0, 0), (-1, -1), "CENTER"), ("TOPPADDING", (0, 0), (-1, 0), 10),
        ("BOTTOMPADDING", (0, 1), (-1, 1), 8), ("LINEBELOW", (0, 1), (-1, 1), 0.5, colors.HexColor("#E5E7EB")),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
    ]))
    story.append(snaptbl)

    # ── Advise the procurement agency ──
    story.append(Paragraph("Advise the procurement agency", h2))
    if c["actions"]:
        for i, a in enumerate(c["actions"], 1):
            if a.startswith("MANDATORY"):
                text = f"<b>{i}.</b> <font color='#B42318'><b>{a}</b></font>"
            else:
                text = f"<b>{i}.</b> {a}"
            story.append(Paragraph(text, body))
            story.append(Spacer(1, 2))
    else:
        story.append(Paragraph("No blocking actions — the specification is ready to tender.", body))

    # ── Key standards to require ──
    story.append(Paragraph("Key standards to require", h2))
    rows = [["Standard", "What to require", "Match"]]
    for r in c["top_rows"]:
        flag = " <b>[QCO MANDATORY]</b>" if r["mandatory"] else ""
        demo = " (DEMO)" if r["is_demo"] else ""
        left = f"<b>{r['is_number']}</b>{flag}{demo}<br/><font size=7 color='#667085'>{r['title']}</font>"
        rows.append([Paragraph(left, body), Paragraph(r["advice"], body),
                     Paragraph(f"{r['match_pct']}%<br/><font size=7 color='#667085'>{r['relevance']}</font>", body)])
    if len(rows) == 1:
        rows.append([Paragraph("No strong matches — see the app for partial matches.", body), "", ""])
    ktbl = Table(rows, colWidths=[doc.width * 0.34, doc.width * 0.52, doc.width * 0.14], repeatRows=1)
    ktbl.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), green), ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTSIZE", (0, 0), (-1, 0), 8.5), ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#E5E7EB")), ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("TOPPADDING", (0, 0), (-1, -1), 5), ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
        ("LEFTPADDING", (0, 0), (-1, -1), 6), ("RIGHTPADDING", (0, 0), (-1, -1), 6),
        ("ALIGN", (2, 0), (2, -1), "CENTER"),
    ]))
    story.append(ktbl)

    if c["appendix_count"] > 0:
        story.append(Spacer(1, 4))
        story.append(Paragraph(
            f"<i>{c['appendix_count']} additional partial/low-relevance match(es) are recorded in MORPHEUS "
            "and omitted from this summary. Open the analysis to review them.</i>", sub))

    # ── Sign-off ──
    story.append(Paragraph("Officer sign-off", h2))
    signoff = [[Paragraph("<b>Reviewed by</b>", body), Paragraph(getattr(officer, "full_name", None) or "____________________", body)],
               [Paragraph("<b>Designation</b>", body), Paragraph((getattr(officer, "role", "") or "____________________").title(), body)],
               [Paragraph("<b>Date</b>", body), Paragraph(datetime.now(timezone.utc).strftime("%d %b %Y"), body)],
               [Paragraph("<b>Signature</b>", body), Paragraph("____________________", body)]]
    sotbl = Table(signoff, colWidths=[doc.width * 0.25, doc.width * 0.75])
    sotbl.setStyle(TableStyle([("LINEBELOW", (1, 0), (1, -1), 0.4, colors.HexColor("#E5E7EB")),
                               ("TOPPADDING", (0, 0), (-1, -1), 6), ("BOTTOMPADDING", (0, 0), (-1, -1), 6)]))
    story.append(sotbl)

    story.append(Spacer(1, 10))
    disc = "This summary records AI-assisted findings for officer verification. Regulatory status is shown only " \
           "from stored records or abstained. "
    if data["has_demo"]:
        disc += "Some standards are DEMO / SYNTHETIC and are not official BIS data. "
    disc += "Prototype — not an official Government of India system."
    story.append(Paragraph(disc, ParagraphStyle("disc", parent=sub, fontSize=7.5)))

    doc.build(story)
    return buf.getvalue()


def build_docx(db: Session, analysis_id: str, officer=None) -> bytes:
    import docx

    data = _gather(db, analysis_id)
    c = _curate(db, analysis_id, data)
    doc_name = data["document"].filename if data["document"] else "—"

    d = docx.Document()
    d.add_heading("Procurement Compliance Summary", level=0)
    d.add_paragraph("MORPHEUS — Standards Intelligence for Public Procurement")
    d.add_paragraph(f"Tender: {data['analysis'].title or doc_name}  |  Sector: {data['analysis'].sector or '—'}  |  "
                    f"Generated: {datetime.now(timezone.utc).strftime('%d %b %Y, %H:%M UTC')}")

    d.add_heading(c["verdict"], level=1)
    d.add_paragraph(c["verdict_detail"])
    d.add_paragraph(f"Compliance {c['compliance_pct']}%  ·  {c['covered']}/{c['requirements_total']} covered  ·  "
                    f"{c['partial']} partial  ·  {c['missing']} not covered  ·  {c['mandatory_count']} QCO-mandatory")

    d.add_heading("Advise the procurement agency", level=1)
    if c["actions"]:
        for a in c["actions"]:
            d.add_paragraph(a, style="List Number")
    else:
        d.add_paragraph("No blocking actions — the specification is ready to tender.")

    d.add_heading("Key standards to require", level=1)
    t = d.add_table(rows=1, cols=3)
    for i, h in enumerate(["Standard", "What to require", "Match"]):
        t.rows[0].cells[i].text = h
    for r in c["top_rows"]:
        cells = t.add_row().cells
        flag = " [QCO MANDATORY]" if r["mandatory"] else ""
        demo = " (DEMO)" if r["is_demo"] else ""
        cells[0].text = f"{r['is_number']}{flag}{demo}\n{r['title']}"
        cells[1].text = r["advice"]
        cells[2].text = f"{r['match_pct']}% {r['relevance']}"
    if c["appendix_count"] > 0:
        d.add_paragraph(f"{c['appendix_count']} additional partial/low-relevance match(es) are recorded in MORPHEUS "
                        "and omitted from this summary.")

    d.add_heading("Officer sign-off", level=1)
    d.add_paragraph(f"Reviewed by: {getattr(officer, 'full_name', None) or '____________________'}")
    d.add_paragraph(f"Designation: {(getattr(officer, 'role', '') or '____________________').title()}")
    d.add_paragraph(f"Date: {datetime.now(timezone.utc).strftime('%d %b %Y')}")
    d.add_paragraph("Signature: ____________________")

    disc = "This summary records AI-assisted findings for officer verification. "
    if data["has_demo"]:
        disc += "Some standards are DEMO / SYNTHETIC and are not official BIS data. "
    disc += "Prototype — not an official Government of India system."
    d.add_paragraph(disc)

    out = io.BytesIO()
    d.save(out)
    return out.getvalue()


def _evidence_for(db: Session, recommendation_id: str):
    ids = [re.evidence_id for re in db.execute(
        select(RecommendationEvidence).where(RecommendationEvidence.recommendation_id == recommendation_id)).scalars()]
    if not ids:
        return []
    return list(db.execute(select(Evidence).where(Evidence.id.in_(ids))).scalars())
