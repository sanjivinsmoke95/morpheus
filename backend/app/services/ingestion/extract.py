"""Text extraction from uploaded documents (ai-pipeline.md §3.1).

Deterministic. PDF via PyMuPDF, DOCX via python-docx, TXT directly. A PDF page
with almost no embedded text is treated as scanned and sent to Tesseract OCR
(when available); OCR confidence is recorded. Nothing here calls an LLM.
"""

from __future__ import annotations

import io
import logging
from dataclasses import dataclass, field

logger = logging.getLogger(__name__)

_SCANNED_CHAR_THRESHOLD = 20  # a page with fewer real chars is likely an image


@dataclass
class ExtractedPage:
    page_number: int
    text: str
    ocr_used: bool = False
    ocr_confidence: float | None = None
    layout: dict = field(default_factory=dict)


@dataclass
class ExtractionResult:
    pages: list[ExtractedPage]

    @property
    def is_scanned(self) -> bool:
        return any(p.ocr_used for p in self.pages)

    @property
    def full_text(self) -> str:
        return "\n".join(p.text for p in self.pages)


def extract(data: bytes, mime_type: str, filename: str) -> ExtractionResult:
    name = (filename or "").lower()
    if mime_type == "application/pdf" or name.endswith(".pdf"):
        return _extract_pdf(data)
    if "wordprocessingml" in mime_type or name.endswith(".docx"):
        return _extract_docx(data)
    return _extract_txt(data)


def _extract_pdf(data: bytes) -> ExtractionResult:
    import fitz  # PyMuPDF

    pages: list[ExtractedPage] = []
    with fitz.open(stream=data, filetype="pdf") as doc:
        for i, page in enumerate(doc, start=1):
            text = page.get_text("text") or ""
            if len(text.strip()) >= _SCANNED_CHAR_THRESHOLD:
                pages.append(ExtractedPage(page_number=i, text=text.strip()))
                continue
            ocr_text, conf = _ocr_pdf_page(page)
            pages.append(
                ExtractedPage(page_number=i, text=ocr_text.strip(), ocr_used=bool(ocr_text.strip()),
                              ocr_confidence=conf, layout={"scanned": True})
            )
    return ExtractionResult(pages=pages)


def _ocr_pdf_page(page) -> tuple[str, float | None]:
    """Render a PDF page and OCR it. Degrades gracefully if Tesseract is absent."""
    try:
        import pytesseract
        from PIL import Image

        pix = page.get_pixmap(dpi=200)
        img = Image.open(io.BytesIO(pix.tobytes("png")))
        text = pytesseract.image_to_string(img)
        data = pytesseract.image_to_data(img, output_type=pytesseract.Output.DICT)
        confs = [int(c) for c in data.get("conf", []) if str(c).lstrip("-").isdigit() and int(c) >= 0]
        conf = round(sum(confs) / len(confs) / 100.0, 3) if confs else None
        return text, conf
    except Exception as exc:  # noqa: BLE001 — OCR unavailable → empty, flagged for review
        logger.warning("OCR unavailable/failed: %s", exc)
        return "", None


def _extract_docx(data: bytes) -> ExtractionResult:
    import docx

    document = docx.Document(io.BytesIO(data))
    parts = [p.text for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    # DOCX has no intrinsic pages; treat as a single logical page.
    return ExtractionResult(pages=[ExtractedPage(page_number=1, text="\n".join(parts))])


def _extract_txt(data: bytes) -> ExtractionResult:
    text = data.decode("utf-8", errors="replace")
    return ExtractionResult(pages=[ExtractedPage(page_number=1, text=text)])
